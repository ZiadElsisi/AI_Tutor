from backend.ingestion.documents import SlideText, chunk_slides, normalize, extract_document
from backend.schemas import QuizQuestion, QuizResponse
from backend.workflows.evaluation import evaluate
from backend.workflows.lecture import analyze
from backend.workflows.quiz import generate, youtube_search_query
from backend.retrieval.index import LectureIndex
from backend import api
from fastapi import HTTPException
from io import BytesIO
import fitz
from pptx import Presentation
from pptx.util import Inches
import asyncio
import httpx


class TinyEmbedder:
    def encode(self, texts, normalize_embeddings=True):
        import numpy as np
        vectors=[]
        for text in texts:
            t=text.lower()
            vectors.append([t.count("gravity")+1, t.count("cell")+1, len(t.split())/100])
        a=np.asarray(vectors,dtype=float)
        if normalize_embeddings:
            a=a/(np.linalg.norm(a,axis=1,keepdims=True)+1e-9)
        return a


class StubRuntime:
    def __init__(self, responses):
        self.responses=iter(responses)
        self.calls=0
    def generate(self, system, user, max_new_tokens=700):
        self.calls+=1
        return next(self.responses)


def test_chunk_preserves_slide_and_overlaps():
    slides=[SlideText(7,"one two three four five six seven eight")]
    chunks=chunk_slides(slides,words=5,overlap=2)
    assert chunks[0]["slide"]==7
    assert chunks[1]["text"].split()[:2]==["four","five"]


def test_retrieval_and_lecture_index_isolation():
    first=LectureIndex([{"text":"Gravity pulls objects toward Earth","slide":1}],TinyEmbedder())
    second=LectureIndex([{"text":"Cells divide and grow","slide":2}],TinyEmbedder())
    assert first.search("gravity")[0]["slide"]==1
    assert second.search("cell")[0]["slide"]==2
    assert all(p["slide"]==1 for p in first.search("cell"))


def test_quiz_scoring_is_deterministic():
    question=QuizQuestion(question="What?",options=["A","B"],correct_index=1,explanation="B.",difficulty="easy",topic="Cells",slide_refs=[3])
    result=evaluate([question],[0])
    assert (result.score,result.percent)==(0,0)
    assert result.suggested_review_topics==["Cells"]
    assert result.feedback[0].slide_refs==[3]


def test_normalize_whitespace():
    assert normalize("  slide\n  text  ")=="slide text"


def test_pdf_extraction_preserves_page_numbers():
    doc=fitz.open()
    page=doc.new_page()
    page.insert_text((72,72),"Newton's laws and gravity")
    pages=extract_document("sample.pdf",doc.tobytes())
    assert len(pages)==1 and pages[0].number==1
    assert "gravity" in pages[0].text


def test_pptx_extraction_preserves_slide_numbers():
    deck=Presentation()
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1))
    box.text="Cell structure and function"
    output=BytesIO()
    deck.save(output)
    slides=extract_document("sample.pptx",output.getvalue())
    assert slides[0].number==1 and "Cell structure" in slides[0].text


def test_invalid_structured_quiz_answer_key_is_rejected():
    import pytest
    with pytest.raises(Exception):
        QuizQuestion(question="Q",options=["A","B"],correct_index=2,explanation="x",difficulty="easy",topic="T")
    with pytest.raises(Exception):
        QuizResponse.model_validate({"quiz_id":"q","questions":[],"unexpected":"untrusted"})


def test_quiz_workflow_retries_and_filters_unverified_slides():
    raw='{"questions":[{"question":"Which force?","options":["gravity","friction"],"correct_index":0,"explanation":"Slide explains gravity.","difficulty":"easy","topic":"Forces","slide_refs":[7,99]}]}'
    runtime=StubRuntime(["not json",raw])
    questions=generate(1,"easy","Forces",[{"text":"Gravity pulls objects","slide":7}],runtime)
    assert runtime.calls==2
    assert questions[0].slide_refs==[7]
    assert youtube_search_query("Forces","Physics")=="Forces Physics lecture explanation"


def test_lecture_analysis_retries_and_filters_slide_references():
    raw='{"title":"Forces","subject":"Physics","topics":[{"name":"Gravity","subtopics":["Mass"],"slide_refs":[1,100]}],"learning_objectives":["Explain gravity"]}'
    runtime=StubRuntime(["invalid",raw])
    result=analyze("Forces","Physics",[{"text":"Gravity pulls objects","slide":1}],2,"slides.pdf","lecture-1",runtime)
    assert runtime.calls==2
    assert result.topics[0].slide_refs==[1]


def test_api_authentication_rejects_missing_and_wrong_secret(monkeypatch):
    import pytest
    monkeypatch.setenv("BACKEND_SECRET","long-secret-value")
    with pytest.raises(HTTPException) as missing:
        api.authenticate(None)
    assert missing.value.status_code==401
    with pytest.raises(HTTPException):
        api.authenticate("wrong")
    api.authenticate("long-secret-value")


def test_unloaded_model_is_reported_without_fake_generation(monkeypatch):
    import pytest
    monkeypatch.setattr(api.runtime,"generator",None)
    monkeypatch.setattr(api.runtime,"embedder",None)
    assert api.health()["model_loaded"] is False
    with pytest.raises(HTTPException) as unavailable:
        api.require_model()
    assert unavailable.value.status_code==503


def test_api_auth_health_missing_model_and_quiz_answer_key(monkeypatch):
    monkeypatch.setenv("BACKEND_SECRET","long-secret-value")
    api.lectures.clear()
    api.quizzes.clear()
    api.runtime.generator=None
    api.runtime.embedder=None
    question=QuizQuestion(question="Q",options=["A","B"],correct_index=1,explanation="B",difficulty="easy",topic="T",slide_refs=[1])
    api.lectures["lecture-1"]={}
    api.quizzes["quiz-1"]={"lecture_id":"lecture-1","questions":[question]}

    async def requests():
        transport=httpx.ASGITransport(app=api.app)
        async with httpx.AsyncClient(transport=transport,base_url="http://test") as client:
            health=await client.get("/health")
            assert health.status_code==200 and health.json()["model_loaded"] is False
            unauthorized=await client.get("/lectures/missing/topics")
            assert unauthorized.status_code==401
            unavailable=await client.post("/lectures/lecture-1/ask",headers={"X-API-Key":"long-secret-value"},json={"question":"Q?"})
            assert unavailable.status_code==503
            scored=await client.post("/lectures/lecture-1/evaluate",headers={"X-API-Key":"long-secret-value"},json={"quiz_id":"quiz-1","answers":[0]})
            assert scored.status_code==200 and scored.json()["score"]==0
            tampered=await client.post("/lectures/lecture-1/evaluate",headers={"X-API-Key":"long-secret-value"},json={"quiz_id":"quiz-1","answers":[1],"questions":[]})
            assert tampered.status_code==422
    asyncio.run(requests())
