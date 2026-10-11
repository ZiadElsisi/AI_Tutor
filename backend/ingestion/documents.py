from dataclasses import dataclass
from io import BytesIO
import re
import fitz
from pptx import Presentation


@dataclass
class SlideText:
    number: int
    text: str


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def extract_document(filename: str, data: bytes) -> list[SlideText]:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        document = fitz.open(stream=data, filetype="pdf")
        return [SlideText(i + 1, normalize(page.get_text())) for i, page in enumerate(document)]
    if lower.endswith(".pptx"):
        deck = Presentation(BytesIO(data))
        slides = []
        for i, slide in enumerate(deck.slides, 1):
            parts = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
            slides.append(SlideText(i, normalize(" ".join(parts))))
        return slides
    raise ValueError("Only PDF and PPTX files are supported.")


def chunk_slides(slides: list[SlideText], words: int = 180, overlap: int = 35) -> list[dict]:
    chunks = []
    for slide in slides:
        tokens = slide.text.split()
        if not tokens:
            continue
        start = 0
        while start < len(tokens):
            part = " ".join(tokens[start:start + words])
            chunks.append({"text": part, "slide": slide.number})
            if start + words >= len(tokens):
                break
            start += max(1, words - overlap)
    return chunks
